"""导出：Markdown / HTML / DOCX。

导出文件同样带着产品的排印性格（宋体正文、等宽数字、发丝表格线），
因为报告最常见的归宿是被打印、被贴进周报、被归档。
"""

from __future__ import annotations

import io
import re
from typing import Any

import markdown as markdown_lib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

_HTML_TEMPLATE = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{
    --ink: #14171a;
    --ink-soft: #5a6068;
    --rule: #d3d6d2;
    --rule-soft: #e6e8e4;
    --face: #f4f4f1;
    --accent: #17787f;
    --pass: #2e7d53;
    --fail: #b02a20;
    --blocked: #b57a12;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    padding: 48px 24px 72px;
    background: var(--face);
    color: var(--ink);
    font-family: "Songti SC", "SimSun", "Source Han Serif SC", "Noto Serif SC", serif;
    font-size: 15px;
    line-height: 1.85;
  }}
  main {{ max-width: 820px; margin: 0 auto; }}
  header {{ border-bottom: 2px solid var(--ink); padding-bottom: 16px; margin-bottom: 28px; }}
  .kicker {{
    font-family: "Microsoft YaHei", "PingFang SC", "Source Han Sans SC", sans-serif;
    font-size: 11px; letter-spacing: .18em; color: var(--accent); margin: 0 0 8px;
  }}
  h1 {{ font-size: 26px; margin: 0 0 6px; letter-spacing: -.01em; }}
  .meta {{
    font-family: "Microsoft YaHei", "PingFang SC", sans-serif;
    font-size: 12px; color: var(--ink-soft); margin: 0;
  }}
  h2 {{
    font-family: "Microsoft YaHei", "PingFang SC", "Source Han Sans SC", sans-serif;
    font-size: 15px; letter-spacing: .04em; margin: 34px 0 10px;
    padding-bottom: 6px; border-bottom: 1px solid var(--rule);
  }}
  h3 {{ font-family: "Microsoft YaHei", sans-serif; font-size: 14px; margin: 22px 0 8px; }}
  p, li {{ margin: 0 0 8px; }}
  ul, ol {{ padding-left: 22px; }}
  code {{
    font-family: "JetBrains Mono", "Cascadia Mono", Consolas, monospace;
    font-size: .88em; background: var(--rule-soft); padding: 1px 5px; border-radius: 2px;
  }}
  pre {{ background: #fff; border: 1px solid var(--rule); padding: 14px 16px; overflow-x: auto; }}
  pre code {{ background: none; padding: 0; }}
  table {{ border-collapse: collapse; width: 100%; margin: 14px 0; font-size: 14px; }}
  th, td {{ border: 1px solid var(--rule); padding: 8px 12px; text-align: left; }}
  th {{
    font-family: "Microsoft YaHei", "PingFang SC", sans-serif;
    font-size: 12px; letter-spacing: .04em; background: #fff;
  }}
  td.num {{ text-align: right; font-family: "JetBrains Mono", Consolas, monospace; }}
  blockquote {{ margin: 14px 0; padding: 2px 0 2px 16px; border-left: 1px solid var(--rule); color: var(--ink-soft); }}
  strong {{ font-weight: 700; }}
  footer {{
    margin-top: 40px; padding-top: 14px; border-top: 1px solid var(--rule);
    font-family: "Microsoft YaHei", sans-serif; font-size: 11px; color: var(--ink-soft);
  }}
  @media print {{
    body {{ background: #fff; padding: 0; }}
    header {{ border-bottom-color: #000; }}
  }}
</style>
</head>
<body>
<main>
<header>
  <p class="kicker">AI 测试报告助手 · 导出存档</p>
  <h1>{title}</h1>
  <p class="meta">{meta}</p>
</header>
{body}
<footer>本报告由 AI 测试报告助手生成 · 指标由系统确定性计算 · 分析文字由 {model} 生成</footer>
</main>
</body>
</html>
"""


def _meta_line(record: dict[str, Any]) -> str:
    parts = []
    if record.get("version"):
        parts.append(f"受检版本 {record['version']}")
    if record.get("created_at"):
        parts.append(f"生成时间 {record['created_at']}")
    if record.get("model"):
        parts.append(f"执行模型 {record['model']}")
    metrics = record.get("metrics") or {}
    if metrics.get("verdict_label"):
        parts.append(f"判定 {metrics['verdict_label']}")
    return " · ".join(parts)


def _metrics_table(record: dict[str, Any]) -> str:
    metrics = record.get("metrics")
    if not metrics:
        return ""
    rows = [
        ("用例总数", metrics.get("total", 0)),
        ("通过", metrics.get("passed", 0)),
        ("失败", metrics.get("failed", 0)),
        ("阻塞", metrics.get("blocked", 0)),
        ("未执行", metrics.get("skipped", 0)),
        ("通过率", f"{metrics.get('pass_rate', 0) * 100:.1f}%"),
        ("执行率", f"{metrics.get('execution_rate', 0) * 100:.1f}%"),
        ("缺陷总数 / 未关闭", f"{metrics.get('defects_total', 0)} / {metrics.get('defects_open', 0)}"),
        ("系统判定", metrics.get("verdict_label", "无法判定")),
    ]
    body = "\n".join(
        f"<tr><th>{label}</th><td class=\"num\">{value}</td></tr>" for label, value in rows
    )
    reasons = metrics.get("verdict_reasons") or []
    reason_html = ""
    if reasons:
        reason_html = "<p><strong>判定依据：</strong>" + "；".join(reasons) + "</p>"
    return f"<h2>检验项目</h2>\n<table>\n{body}\n</table>\n{reason_html}"


def to_markdown(record: dict[str, Any]) -> str:
    lines = [f"# {record.get('title') or '测试报告'}", ""]
    meta = _meta_line(record)
    if meta:
        lines += [meta, ""]
    metrics = record.get("metrics")
    if metrics:
        lines += [
            "## 检验项目",
            "",
            "| 项目 | 数值 |",
            "| --- | --- |",
            f"| 用例总数 | {metrics.get('total', 0)} |",
            f"| 通过 | {metrics.get('passed', 0)} |",
            f"| 失败 | {metrics.get('failed', 0)} |",
            f"| 阻塞 | {metrics.get('blocked', 0)} |",
            f"| 未执行 | {metrics.get('skipped', 0)} |",
            f"| 通过率 | {metrics.get('pass_rate', 0) * 100:.1f}% |",
            f"| 执行率 | {metrics.get('execution_rate', 0) * 100:.1f}% |",
            f"| 缺陷总数 / 未关闭 | {metrics.get('defects_total', 0)} / {metrics.get('defects_open', 0)} |",
            f"| 系统判定 | {metrics.get('verdict_label', '无法判定')} |",
            "",
        ]
        reasons = metrics.get("verdict_reasons") or []
        if reasons:
            lines += [f"**判定依据：**{'；'.join(reasons)}", ""]
    lines += [record.get("report") or "", ""]
    lines += [
        "---",
        f"本报告由 AI 测试报告助手生成；指标由系统确定性计算，分析文字由 {record.get('model', '—')} 生成。",
    ]
    return "\n".join(lines)


def to_html(record: dict[str, Any]) -> str:
    body_md = (record.get("report") or "").strip()
    body_html = markdown_lib.markdown(
        body_md, extensions=["tables", "fenced_code", "sane_lists", "nl2br"]
    )
    return _HTML_TEMPLATE.format(
        title=record.get("title") or "测试报告",
        meta=_meta_line(record),
        model=record.get("model") or "—",
        body=f"{_metrics_table(record)}\n{body_html}",
    )


# ------------------------------------------------------------------ DOCX

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_TABLE_SEP_RE = re.compile(r"^\|[\s:\-|]+\|$")


def _add_runs(paragraph, text: str) -> None:
    text = text.replace("`", "")
    position = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _write_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            buffer: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                buffer.append(lines[index])
                index += 1
            index += 1
            paragraph = doc.add_paragraph()
            run = paragraph.add_run("\n".join(buffer))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and _TABLE_SEP_RE.match(lines[index + 1].strip()):
            header = _split_row(stripped)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index]))
                index += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, header):
                cell.text = value.replace("**", "")
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value.replace("**", "")
            doc.add_paragraph()
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            _add_runs(paragraph, stripped[2:])
        elif stripped:
            _add_runs(doc.add_paragraph(), stripped)
        index += 1


def to_docx(record: dict[str, Any]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)

    heading = doc.add_heading(record.get("title") or "测试报告", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = _meta_line(record)
    if meta:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(meta)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x5A, 0x60, 0x68)

    metrics = record.get("metrics")
    if metrics:
        doc.add_heading("检验项目", level=2)
        rows = [
            ("用例总数", metrics.get("total", 0)),
            ("通过", metrics.get("passed", 0)),
            ("失败", metrics.get("failed", 0)),
            ("阻塞", metrics.get("blocked", 0)),
            ("未执行", metrics.get("skipped", 0)),
            ("通过率", f"{metrics.get('pass_rate', 0) * 100:.1f}%"),
            ("执行率", f"{metrics.get('execution_rate', 0) * 100:.1f}%"),
            ("缺陷总数 / 未关闭", f"{metrics.get('defects_total', 0)} / {metrics.get('defects_open', 0)}"),
            ("系统判定", metrics.get("verdict_label", "无法判定")),
        ]
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = str(label)
            cells[1].text = str(value)
        reasons = metrics.get("verdict_reasons") or []
        if reasons:
            _add_runs(doc.add_paragraph(), "**判定依据：**" + "；".join(reasons))

    defects = record.get("defects") or []
    if defects:
        doc.add_heading("缺陷清单", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("编号", "优先级", "模块", "状态")):
            cell.text = value
        for defect in defects:
            cells = table.add_row().cells
            cells[0].text = str(defect.get("id") or "—")
            cells[1].text = str(defect.get("priority") or "—")
            cells[2].text = str(defect.get("module") or "—")
            cells[3].text = str(defect.get("status") or "—")

    _write_markdown(doc, record.get("report") or "")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
